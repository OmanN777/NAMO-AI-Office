import os
import openpyxl
from copy import copy

workspace_dir = r"C:\Users\namo_\OneDrive\เอกสาร\gemini-cli\antigravity-office-workspace"
backup_path = os.path.join(workspace_dir, "Work_Brief", "TrueMoney_QA_Test_Answers", "facebook_login_testcases.xlsx")
target_path = os.path.join(workspace_dir, "Work_Brief", "TrueMoney_QA_Test_Answers", "solution2_facebook_login_testcases.xlsx")

def update_boss_excel():
    wb = openpyxl.load_workbook(backup_path)
    ws = wb.active

    # New test cases requested by Boss to be appended
    new_rows = [
        [
            "Blank Validation - Email Blank",
            "To verify inline validation when leaving email/phone field blank while entering password.",
            "1. Facebook Mobile App installed.",
            "1. Open App.\n2. Leave Email/Phone field empty.\n3. Enter valid password.\n4. Tap 'Log In' button.",
            "1. Login action prevented.\n2. Inline error displayed: 'Please enter your email address or phone number'.",
            "High"
        ],
        [
            "Blank Validation - Password Blank",
            "To verify inline validation when entering valid email/phone while leaving password field blank.",
            "1. Facebook Mobile App installed.",
            "1. Open App.\n2. Enter valid registered email.\n3. Leave Password field empty.\n4. Tap 'Log In' button.",
            "1. Login action prevented.\n2. Inline error displayed: 'Please enter your password'.",
            "High"
        ],
        [
            "Invalid Email Format Validation",
            "To verify app validation when user enters email with invalid format (missing '@' or domain).",
            "1. Facebook Mobile App installed.",
            "1. Open App.\n2. Enter invalid email format (e.g. 'namo.nanon5gmail.com' or 'namo@').\n3. Enter any password.\n4. Tap 'Log In'.",
            "1. Login action prevented.\n2. Validation message displayed: 'Please enter a valid email address'.",
            "Medium"
        ],
        [
            "Trailing & Leading Space in Email",
            "To verify that app automatically trims leading/trailing whitespaces in email input field.",
            "1. Facebook Mobile App installed.\n2. Active internet connection.",
            "1. Open App.\n2. Enter email with spaces: '  namo.nanon5@gmail.com  '.\n3. Enter correct password.\n4. Tap 'Log In'.",
            "1. App automatically trims whitespace.\n2. User logs in successfully to News Feed.",
            "Medium"
        ],
        [
            "Password Case Sensitivity Check",
            "To verify that password field is strictly case-sensitive.",
            "1. Facebook Mobile App installed.\n2. Active internet connection.\n3. Valid account password: 'Password123!'.",
            "1. Open App.\n2. Enter valid email.\n3. Enter password in wrong case: 'password123!'.\n4. Tap 'Log In'.",
            "1. Login fails.\n2. Error message displayed: 'Incorrect Password' (verifying case-sensitivity).",
            "High"
        ]
    ]

    # Reference template cell formatting from existing row (e.g. Row 2)
    template_row = 2

    start_row = ws.max_row + 1
    for i, row_data in enumerate(new_rows):
        current_row = start_row + i
        ws.row_dimensions[current_row].height = 65

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            
            # Copy style from Boss's template row
            ref_cell = ws.cell(row=template_row, column=col_idx)
            if ref_cell.font:
                cell.font = copy(ref_cell.font)
            if ref_cell.alignment:
                cell.alignment = copy(ref_cell.alignment)
            if ref_cell.border:
                cell.border = copy(ref_cell.border)
            if ref_cell.fill:
                cell.fill = copy(ref_cell.fill)

            # Match priority formatting for col 6
            if col_idx == 6:
                p_val = str(val).strip().lower()
                for r in range(2, ws.max_row):
                    match_cell = ws.cell(row=r, column=6)
                    if match_cell.value and str(match_cell.value).strip().lower() == p_val:
                        if match_cell.fill:
                            cell.fill = copy(match_cell.fill)
                        if match_cell.font:
                            cell.font = copy(match_cell.font)
                        break

    try:
        wb.save(target_path)
        print("Successfully updated target Excel file.")
    except PermissionError:
        alt_target = os.path.join(workspace_dir, "Work_Brief", "TrueMoney_QA_Test_Answers", "solution2_facebook_login_testcases_updated.xlsx")
        wb.save(alt_target)
        print(f"Target file was open in Excel. Saved updated copy to solution2_facebook_login_testcases_updated.xlsx")

    try:
        wb.save(backup_path)
        print("Successfully updated backup Excel file.")
    except PermissionError:
        print("Backup file is currently open in Excel. Please close it if you want to overwrite it directly.")

if __name__ == "__main__":
    update_boss_excel()
