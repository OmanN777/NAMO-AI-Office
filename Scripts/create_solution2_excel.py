import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

workspace_dir = r"C:\Users\namo_\OneDrive\เอกสาร\gemini-cli\antigravity-office-workspace"
output_path = os.path.join(workspace_dir, "Work_Brief", "TrueMoney_QA_Test_Answers", "solution2_facebook_login_testcases.xlsx")

def create_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Facebook Login Test Cases"

    # Ensure gridlines are visible
    ws.views.sheetView[0].showGridLines = True

    # Table Headers matching exact TrueMoney PDF column names
    headers = [
        "Test case name",
        "Objective",
        "Prerequisite",
        "Test Steps",
        "Expected Result",
        "Priority"
    ]

    # Enhanced Data Rows including requested test cases:
    # 1. Blank Email & Blank Password separated
    # 2. Invalid Email Format
    # 3. Trailing/Leading Space in Email
    # 4. Password Case Sensitivity
    rows = [
        [
            "Valid Login with Email",
            "To verify that user can login successfully when they put a correct registered email and password.",
            "1. Facebook Mobile App installed.\n2. Active internet connection.\n3. Registered account with email.",
            "1. Open Facebook Mobile App.\n2. Enter registered email.\n3. Enter correct password.\n4. Tap 'Log In' button.",
            "1. Login successful.\n2. Redirected to Facebook News Feed.\n3. User session token saved.",
            "High"
        ],
        [
            "Valid Login with Phone Number",
            "To verify that user can login successfully when they put a correct registered phone number and password.",
            "1. Facebook Mobile App installed.\n2. Active internet connection.\n3. Registered account with phone number.",
            "1. Open Facebook Mobile App.\n2. Enter registered phone number (e.g. 0812345678).\n3. Enter correct password.\n4. Tap 'Log In' button.",
            "1. Login successful.\n2. Redirected to Facebook News Feed.",
            "High"
        ],
        [
            "Invalid Login - Password Incorrect",
            "To verify that user can login unsuccessfully when they put a correct username but wrong password.",
            "1. Facebook Mobile App installed.\n2. Active internet connection.",
            "1. Open Facebook Mobile App.\n2. Enter valid email/phone.\n3. Enter incorrect password.\n4. Tap 'Log In' button.",
            "1. Login failed.\n2. Alert message displayed: 'Incorrect Password'.\n3. Password field cleared.",
            "High"
        ],
        [
            "Invalid Login - Username Not Found",
            "To verify that user can login unsuccessfully when they put an unregistered email or phone number.",
            "1. Facebook Mobile App installed.\n2. Active internet connection.",
            "1. Open Facebook Mobile App.\n2. Enter non-existent email (e.g. unknown_user9999@test.com).\n3. Enter any password.\n4. Tap 'Log In' button.",
            "1. Login failed.\n2. Alert message displayed: 'Invalid Credentials' or 'Account Not Found'.",
            "High"
        ],
        [
            "Blank Validation - Email Blank",
            "To verify inline validation when leaving email/phone field blank while entering password.",
            "1. Facebook Mobile App installed.",
            "1. Open App.\n2. Leave Email/Phone field empty.\n3. Enter valid password.\n4. Tap 'Log In' button.",
            "1. Login action prevented.\n2. Inline error displayed: 'Please enter your email address or phone number'.",
            "High"
        ],
        [
            "Blank Validation - Password Blank",
            "To verify inline validation when entering valid email/phone while leaving password field blank.",
            "1. Facebook Mobile App installed.",
            "1. Open App.\n2. Enter valid registered email.\n3. Leave Password field empty.\n4. Tap 'Log In' button.",
            "1. Login action prevented.\n2. Inline error displayed: 'Please enter your password'.",
            "High"
        ],
        [
            "Blank Validation - Both Fields Blank",
            "To verify inline validation when tapping Log In with both email/phone and password fields blank.",
            "1. Facebook Mobile App installed.",
            "1. Open App.\n2. Leave both Email/Phone and Password fields empty.\n3. Tap 'Log In' button.",
            "1. Login action prevented.\n2. Prompt message displays: 'Please enter your email and password'.",
            "Medium"
        ],
        [
            "Invalid Email Format Validation",
            "To verify app validation when user enters email with invalid format (missing '@' or domain).",
            "1. Facebook Mobile App installed.",
            "1. Open App.\n2. Enter invalid email format (e.g. 'namo.nanon5gmail.com' or 'namo@').\n3. Enter any password.\n4. Tap 'Log In'.",
            "1. Login action prevented.\n2. Validation message displayed: 'Please enter a valid email address'.",
            "Medium"
        ],
        [
            "Trailing & Leading Space in Email",
            "To verify that app automatically trims leading/trailing whitespaces in email input field.",
            "1. Facebook Mobile App installed.\n2. Active internet connection.",
            "1. Open App.\n2. Enter email with spaces: '  namo.nanon5@gmail.com  '.\n3. Enter correct password.\n4. Tap 'Log In'.",
            "1. App automatically trims whitespace.\n2. User logs in successfully to News Feed.",
            "Medium"
        ],
        [
            "Password Case Sensitivity Check",
            "To verify that password field is strictly case-sensitive.",
            "1. Facebook Mobile App installed.\n2. Active internet connection.\n3. Valid account password: 'Password123!'.",
            "1. Open App.\n2. Enter valid email.\n3. Enter password in wrong case: 'password123!'.\n4. Tap 'Log In'.",
            "1. Login fails.\n2. Error message displayed: 'Incorrect Password' (verifying case-sensitivity).",
            "High"
        ],
        [
            "Forgot Password Flow",
            "To verify password recovery flow via 'Forgot Password?' link.",
            "1. Facebook Mobile App installed.\n2. Active internet connection.",
            "1. Open Facebook Mobile App.\n2. Tap 'Forgot Password?' link.\n3. Enter registered email/phone.\n4. Tap 'Search Account'.",
            "1. User account identified.\n2. Displays options to send OTP code via SMS or Email.",
            "High"
        ],
        [
            "Two-Factor Authentication (2FA)",
            "To verify 2FA prompt when logging in from a new unrecognized mobile device.",
            "1. Account has 2FA enabled.\n2. New unrecognized device.",
            "1. Open Facebook Mobile App.\n2. Enter valid email and password.\n3. Tap 'Log In' button.",
            "1. 2FA verification screen is displayed.\n2. Prompts for 6-digit OTP code sent via SMS or Authenticator App.",
            "High"
        ],
        [
            "Offline Network Handling",
            "To verify app behavior when attempting to login without internet connection.",
            "1. Mobile Wi-Fi and Cellular Data turned OFF.",
            "1. Open Facebook Mobile App.\n2. Enter valid credentials.\n3. Tap 'Log In' button.",
            "1. Error popup displays: 'No Internet Connection. Please check your network and try again.'",
            "High"
        ],
        [
            "Show/Hide Password Toggle",
            "To verify visibility toggle icon in password input field.",
            "1. Facebook Mobile App installed.",
            "1. Open App.\n2. Enter password text.\n3. Tap eye icon (Show password).\n4. Tap eye icon again (Hide password).",
            "1. Password text becomes visible.\n2. Password text gets masked again (••••••).",
            "Low"
        ],
        [
            "App Relaunch Session Persistence",
            "To verify that user remains logged in after closing and reopening the app.",
            "1. User logged in successfully.",
            "1. Close Facebook App from Recent Apps (Kill process).\n2. Relaunch Facebook App.",
            "1. App opens directly to News Feed without requiring re-login.",
            "High"
        ]
    ]

    # Styles
    header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    body_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    high_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    med_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    low_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

    # Write Headers
    ws.append(headers)
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 28

    # Write Data Rows
    for row_idx, row_data in enumerate(rows, start=2):
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 65  # Give height for multi-line text

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = body_font
            cell.border = thin_border
            
            # Alignments
            if col_idx in [1, 6]:  # Test case name & Priority
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

            # Priority Highlights
            if col_idx == 6:
                val = str(cell.value).strip().lower()
                if val == "high":
                    cell.fill = high_fill
                    cell.font = Font(name="Calibri", size=10, bold=True, color="C00000")
                elif val == "medium":
                    cell.fill = med_fill
                    cell.font = Font(name="Calibri", size=10, bold=True, color="B25900")
                elif val == "low":
                    cell.fill = low_fill
                    cell.font = Font(name="Calibri", size=10, bold=True, color="375623")

    # Set Widths
    col_widths = {
        "A": 26, # Test case name
        "B": 32, # Objective
        "C": 30, # Prerequisite
        "D": 35, # Test Steps
        "E": 35, # Expected Result
        "F": 12  # Priority
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(output_path)
    print("Created updated Excel solution file successfully.")

if __name__ == "__main__":
    create_excel()
