import os
import openpyxl
from openpyxl.styles import Alignment

workspace_dir = r"C:\Users\namo_\OneDrive\เอกสาร\gemini-cli\antigravity-office-workspace"
file1 = os.path.join(workspace_dir, "Work_Brief", "TrueMoney_QA_Test_Answers", "facebook_login_testcases.xlsx")
file2 = os.path.join(workspace_dir, "Work_Brief", "TrueMoney_QA_Test_Answers", "solution2_facebook_login_testcases.xlsx")

def fix_wrapping(file_path):
    if not os.path.exists(file_path):
        return
    try:
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active

        # Header Row
        ws.row_dimensions[1].height = 28
        for col_idx in range(1, 7):
            cell = ws.cell(row=1, column=col_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Data Rows
        for r in range(2, ws.max_row + 1):
            ws.row_dimensions[r].height = 70  # Sufficient height for multi-line text
            for c in range(1, 7):
                cell = ws.cell(row=r, column=c)
                if c in [1, 6]:
                    cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        wb.save(file_path)
        print(f"Fixed wrap_text for: {file_path}")
    except PermissionError:
        print(f"Permission denied for {file_path} (File is currently open in Excel).")

if __name__ == "__main__":
    fix_wrapping(file1)
    fix_wrapping(file2)
